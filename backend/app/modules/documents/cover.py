"""Cover letters inherit the uploaded document's page and typography settings.

For exact letter layout use a DOCX letter template containing {{cover_letter}}.
A resume and a prose letter have different body structures; defaults preserve the
resume's title/contact block and theme, not a fictitious identical body layout.
"""

import io
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
from app.modules.documents.docx_writer import _splice_runs
from app.modules.resumes.docx_parser import iter_paragraphs, parse_docx
from app.modules.resumes.latex_parser import escape_latex_text


def save_package(original, document, destination):
    """Only document.xml changes. All other ZIP entry payloads remain identical."""
    xml = etree.tostring(
        document._element, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    Path(destination).parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(io.BytesIO(original)) as src, ZipFile(destination, "w") as dst:
        for entry in src.infolist():
            dst.writestr(
                entry,
                xml
                if entry.filename == "word/document.xml"
                else src.read(entry.filename),
            )


def render_cover_docx(
    original: bytes, text: str, destination: str, template: bytes | None = None
):
    source = template or original
    doc = Document(io.BytesIO(source))
    if template:
        replaced = False
        for p in iter_paragraphs(doc):
            if "{{cover_letter}}" in p.text:
                old = "".join(r.text for r in p.runs)
                if "{{cover_letter}}" not in old:
                    raise ValueError("Cover placeholder must be in ordinary text runs")
                _splice_runs(p.runs, old, old.replace("{{cover_letter}}", text))
                replaced = True
        if not replaced:
            raise ValueError("Cover template must contain {{cover_letter}}")
    else:
        parsed = parse_docx(original)
        # Retain only the title/contact block before the first real section.
        headers = [
            s.locator
            for s in parsed.sections
            if s.locator is not None and s.name != "preamble_body"
        ]
        keep_count = min(headers) if headers else 0
        paragraphs = iter_paragraphs(doc)
        for index, p in reversed(list(enumerate(paragraphs))):
            if index >= keep_count:
                p._element.getparent().remove(p._element)
        # Remove now-empty body tables to avoid retaining resume layout as blank space.
        for table in list(doc.tables):
            if not table._element.xpath(".//w:p"):
                table._element.getparent().remove(table._element)
        # Derive paragraph/run typography without taking bullet numbering.
        sample_doc = Document(io.BytesIO(original))
        sample = next(
            (
                p
                for p in iter_paragraphs(sample_doc)
                if len(p.text) > 40 and "\t" not in p.text
            ),
            None,
        )
        for block in text.strip().split("\n\n"):
            p = doc.add_paragraph()
            if sample is not None:
                if sample._p.pPr is not None:
                    props = deepcopy(sample._p.pPr)
                    for number in props.findall(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr"
                    ):
                        props.remove(number)
                    p._p.insert(0, props)
                if sample.runs and sample.runs[0]._r.rPr is not None:
                    r = p.add_run(block)
                    r._r.insert(0, deepcopy(sample.runs[0]._r.rPr))
                else:
                    p.add_run(block)
            else:
                p.add_run(block)
    save_package(source, doc, destination)
    return destination


def render_cover_tex(original: str, text: str, destination: str):
    begin = original.find(r"\begin{document}")
    if begin < 0:
        raise ValueError("LaTeX resume has no document environment")
    body_start = begin + len(r"\begin{document}")
    import re

    section = re.search(r"\\section\*?\{", original[body_start:])
    header_end = body_start + section.start() if section else body_start
    letter = "\n\n".join(escape_latex_text(p) for p in text.strip().split("\n\n"))
    Path(destination).write_text(
        original[:header_end] + "\n" + letter + "\n\\end{document}\n", encoding="utf-8"
    )
    return destination
