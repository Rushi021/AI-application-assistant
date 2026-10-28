"""Contract and end-to-end approval tests; no employer submissions or vendor calls."""

import base64
import io
from datetime import datetime, timedelta
from pathlib import Path
from zipfile import ZipFile
import httpx
import pytest
from sqlalchemy import select
from app.modules.agent.discovery import JobDiscovery, Posting, canonical
from app.modules.agent.models import AgentPlan, AgentItem
from app.modules.agent.schemas import PlanConfig, Source
from app.modules.agent.service import run_plan
from app.modules.ai.gateway import ModelGateway
from tests.conftest import AUTH, build_sample_docx


class Discovery:
    def search(self, config):
        return [
            Posting(
                "greenhouse",
                "a",
                "Software Engineer",
                "Example",
                "Remote",
                "Required: Python, SQL",
                "https://boards.greenhouse.io/example/jobs/123?utm_source=test",
            )
        ], []


def setup_packet(client):
    resume = client.post(
        "/resumes/upload",
        headers=AUTH,
        data={"family": "ai_ml", "version": "1"},
        files={
            "file": (
                "resume.docx",
                build_sample_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    ).json()
    assert "id" in resume, resume
    client.put(
        "/profile",
        headers=AUTH,
        json={
            "first_name": "Ada",
            "last_name": "Lovelace",
            "email": "ada@example.test",
        },
    )
    config = {
        "enabled": True,
        "resume_id": resume["id"],
        "roles": ["Software Engineer"],
        "locations": ["Remote"],
        "sources": [{"provider": "greenhouse", "board": "example"}],
        "minimum_score": 0,
    }
    response = client.put("/agent/plan", headers=AUTH, json=config)
    assert response.status_code == 200, response.text
    from app.db import session

    with session.SessionLocal() as db:
        plan = db.scalar(select(AgentPlan))
        assert run_plan(db, plan.id, ModelGateway(), discovery=Discovery(), force=True)
    items = client.get("/agent/queue", headers=AUTH).json()
    assert len(items) == 1, items
    assert items[0]["state"] == "ready", items
    return resume, items[0]


def approve(client, item):
    response = client.post(
        f"/agent/queue/{item['id']}/decision", headers=AUTH, json={"action": "approve"}
    )
    assert response.status_code == 200, response.text


def test_prepare_review_approve_claim_upload_and_confirm(client):
    _, item = setup_packet(client)
    # Preparation must never itself produce an application or claimable packet.
    assert client.get("/applications", headers=AUTH).json() == []
    assert client.post("/agent/claim", headers=AUTH, json={}).json() is None
    document = client.get(f"/agent/queue/{item['id']}/documents/resume", headers=AUTH)
    assert document.status_code == 200
    assert base64.b64decode(document.json()["base64"]).startswith(b"PK")
    cover = client.get(
        f"/agent/queue/{item['id']}/documents/cover_letter", headers=AUTH
    ).json()
    assert cover["name"].endswith(".docx")
    approve(client, item)
    packet = client.post("/agent/claim", headers=AUTH, json={}).json()
    assert packet["state"] == "filling"
    assert packet["profile"]["email"] == "ada@example.test"
    assert client.post("/agent/claim", headers=AUTH, json={}).json() is None
    assert (
        client.post(
            f"/agent/queue/{item['id']}/result",
            headers=AUTH,
            json={"lease_token": "wrong", "state": "submitted", "url": packet["url"]},
        ).status_code
        == 409
    )
    response = client.post(
        f"/agent/queue/{item['id']}/result",
        headers=AUTH,
        json={
            "lease_token": packet["lease_token"],
            "state": "submitted",
            "url": packet["url"],
            "detail": "Confirmation detected",
        },
    )
    assert response.json()["state"] == "submitted", response.text
    apps = client.get("/applications", headers=AUTH).json()
    assert len(apps) == 1 and apps[0]["status"] == "applied"
    assert (
        client.post(
            f"/agent/queue/{item['id']}/decision",
            headers=AUTH,
            json={"action": "reprepare"},
        ).status_code
        == 409
    )


def test_document_edit_and_profile_changes_revoke_approval(client):
    _, item = setup_packet(client)
    approve(client, item)
    client.put("/profile", headers=AUTH, json={"first_name": "Changed"})
    assert client.post("/agent/claim", headers=AUTH, json={}).json() is None
    assert client.get("/agent/queue", headers=AUTH).json()[0]["state"] == "ready"
    approve(client, item)
    patch = client.patch(
        f"/tailoring/runs/{item['run_id']}/cover-letter",
        headers=AUTH,
        json={"text": "Updated truthful cover letter."},
    )
    assert patch.status_code == 200
    assert client.post("/agent/claim", headers=AUTH, json={}).json() is None
    assert (
        client.post(
            f"/agent/queue/{item['id']}/decision",
            headers=AUTH,
            json={"action": "approve"},
        ).status_code
        == 409
    )
    client.post(f"/tailoring/runs/{item['run_id']}/finalize", headers=AUTH)
    approve(client, item)


def test_idempotent_discovery_and_uncertain_submission_is_not_retried(client):
    _, item = setup_packet(client)
    from app.db import session

    with session.SessionLocal() as db:
        plan = db.scalar(select(AgentPlan))
        run_plan(db, plan.id, ModelGateway(), discovery=Discovery(), force=True)
    assert len(client.get("/agent/queue", headers=AUTH).json()) == 1
    approve(client, item)
    client.post("/agent/claim", headers=AUTH, json={})
    with session.SessionLocal() as db:
        row = db.get(AgentItem, item["id"])
        row.lease_until = datetime.utcnow() - timedelta(seconds=1)
        db.commit()
    assert client.post("/agent/claim", headers=AUTH, json={}).json() is None
    row = client.get("/agent/queue", headers=AUTH).json()[0]
    assert row["state"] == "needs_attention"
    assert (
        client.post(
            f"/agent/queue/{item['id']}/decision",
            headers=AUTH,
            json={"action": "reprepare"},
        ).status_code
        == 409
    )


def test_dedup_url_safety_and_source_validation():
    assert (
        canonical("https://jobs.example.test/a/?utm_source=mail&job=1#x")
        == "https://jobs.example.test/a?job=1"
    )
    for url in [
        "http://example.com",
        "https://localhost/test",
        "https://127.0.0.1/test",
        "https://169.254.169.254/",
        "https://user:pass@example.com/",
    ]:
        with pytest.raises(ValueError):
            canonical(url)
    with pytest.raises(ValueError):
        Source(provider="greenhouse", board="../../private")


@pytest.mark.parametrize(
    "provider,payload",
    [
        (
            "greenhouse",
            {
                "jobs": [
                    {
                        "id": 1,
                        "title": "Software Engineer",
                        "location": {"name": "Remote"},
                        "content": "<p>Python &amp; SQL</p>",
                        "absolute_url": "https://boards.greenhouse.io/example/jobs/1",
                    }
                ]
            },
        ),
        (
            "lever",
            [
                {
                    "id": "a",
                    "text": "Software Engineer",
                    "categories": {"location": "Remote"},
                    "descriptionPlain": "Python SQL",
                    "lists": [],
                    "hostedUrl": "https://jobs.lever.co/example/a",
                }
            ],
        ),
        (
            "ashby",
            {
                "jobs": [
                    {
                        "id": "a",
                        "title": "Software Engineer",
                        "location": "US",
                        "isRemote": True,
                        "descriptionPlain": "Python SQL",
                        "jobUrl": "https://jobs.ashbyhq.com/example/a",
                    }
                ]
            },
        ),
    ],
)
def test_provider_response_contracts(provider, payload):
    client = httpx.Client(
        transport=httpx.MockTransport(lambda request: httpx.Response(200, json=payload))
    )
    service = JobDiscovery(client)
    config = PlanConfig(
        resume_id=1,
        roles=["Software Engineer"],
        locations=["Remote"],
        sources=[Source(provider=provider, board="example")],
    )
    posts, errors = service.search(config)
    assert len(posts) == 1 and not errors
    assert "Python" in posts[0].description
    service.close()


def test_provider_failure_is_visible_and_other_sources_continue():
    def handler(request):
        if "greenhouse" in request.url.host:
            return httpx.Response(429)
        return httpx.Response(
            200,
            json={
                "jobs": [
                    {
                        "id": "a",
                        "title": "Software Engineer",
                        "location": "Remote",
                        "descriptionPlain": "Python",
                        "jobUrl": "https://jobs.ashbyhq.com/a",
                    }
                ]
            },
        )

    service = JobDiscovery(httpx.Client(transport=httpx.MockTransport(handler)))
    posts, errors = service.search(
        PlanConfig(
            resume_id=1,
            roles=["Software Engineer"],
            locations=["Remote"],
            sources=[
                Source(provider="greenhouse", board="x"),
                Source(provider="ashby", board="x"),
            ],
        )
    )
    assert len(posts) == 1 and len(errors) == 1


def test_no_writes_without_auth(client):
    assert client.post("/agent/claim", json={}).status_code == 401
    assert client.get("/agent/queue").status_code == 401


def test_cover_template_and_docx_package_preservation(client, tmp_path):
    from docx import Document
    from app.modules.documents.docx_writer import write_docx
    from app.modules.documents.cover import render_cover_docx

    original = build_sample_docx()
    out = tmp_path / "tailored.docx"
    write_docx(original, {9: "Built Python pipelines for analytics"}, str(out))
    with ZipFile(io.BytesIO(original)) as a, ZipFile(out) as b:
        assert a.namelist() == b.namelist()
        for name in a.namelist():
            if name != "word/document.xml":
                assert a.read(name) == b.read(name)
    template = Document()
    template.add_paragraph("Ada Lovelace\tEngineering")
    p = template.add_paragraph()
    p.add_run("{{cover_").bold = True
    p.add_run("letter}}")
    buf = io.BytesIO()
    template.save(buf)
    cover = tmp_path / "letter.docx"
    render_cover_docx(
        original,
        "Dear Hiring Manager,\n\nTruthful experience.",
        str(cover),
        template=buf.getvalue(),
    )
    written = Document(cover)
    assert written.paragraphs[0].text == "Ada Lovelace\tEngineering"
    assert "Truthful experience." in written.paragraphs[1].text
    assert written.paragraphs[1].runs[0].bold is True


def test_execution_revalidates_approval_immediately_before_submit(client):
    _, item = setup_packet(client)
    approve(client, item)
    packet = client.post("/agent/claim", headers=AUTH, json={}).json()
    path = f"/agent/queue/{item['id']}/validate"
    assert client.post(
        path, headers=AUTH, json={"lease_token": packet["lease_token"]}
    ).json() == {"valid": True}
    client.put("/profile", headers=AUTH, json={"email": "changed@example.test"})
    assert (
        client.post(
            path, headers=AUTH, json={"lease_token": packet["lease_token"]}
        ).status_code
        == 409
    )


def test_explicit_manual_resolution_requires_new_approval_for_retry(client):
    _, item = setup_packet(client)
    approve(client, item)
    packet = client.post("/agent/claim", headers=AUTH, json={}).json()
    client.post(
        f"/agent/queue/{item['id']}/result",
        headers=AUTH,
        json={
            "lease_token": packet["lease_token"],
            "state": "needs_attention",
            "url": packet["url"],
            "detail": "No confirmation observed",
        },
    )
    response = client.post(
        f"/agent/queue/{item['id']}/resolve",
        headers=AUTH,
        json={"outcome": "not_submitted"},
    )
    assert response.json()["state"] == "ready"
    assert client.post("/agent/claim", headers=AUTH, json={}).json() is None
    approve(client, item)


def test_one_off_scan_does_not_enable_recurring_plan(client):
    setup_packet(client)
    from app.db import session

    with session.SessionLocal() as db:
        plan = db.scalar(select(AgentPlan))
        plan.enabled = False
        plan.run_once = False
        db.commit()
    assert client.post("/agent/scan", headers=AUTH).status_code == 200
    with session.SessionLocal() as db:
        plan = db.scalar(select(AgentPlan))
        assert not plan.enabled and plan.run_once
        assert run_plan(db, plan.id, ModelGateway(), discovery=Discovery())
        db.refresh(plan)
        assert not plan.enabled and not plan.run_once


def test_source_update_revokes_approved_packet(client):
    _, item = setup_packet(client)
    approve(client, item)

    class Changed(Discovery):
        def search(self, config):
            posts, errors = super().search(config)
            posts[0].description = "Required: Python, SQL, Kubernetes"
            return posts, errors

    from app.db import session

    with session.SessionLocal() as db:
        plan = db.scalar(select(AgentPlan))
        run_plan(db, plan.id, ModelGateway(), discovery=Changed(), force=True)
    row = client.get("/agent/queue", headers=AUTH).json()[0]
    assert row["state"] == "ready" and row["run_id"] != item["run_id"]
    assert client.post("/agent/claim", headers=AUTH, json={}).json() is None


def test_manual_tex_round_trip_keeps_exact_source_when_untailored(client):
    from tests.conftest import SAMPLE_TEX

    original = SAMPLE_TEX.replace(
        "\\section{Experience}", "\\section{Experience}  % keep this whitespace\n\n"
    )
    resume = client.post(
        "/resumes/upload",
        headers=AUTH,
        data={"family": "ai_ml", "version": "1"},
        files={"file": ("resume.tex", original.encode(), "text/plain")},
    ).json()
    job = client.post(
        "/jobs/ingest", headers=AUTH, json={"raw_text": "Required: Python"}
    ).json()
    run = client.post(
        "/tailoring/runs",
        headers=AUTH,
        json={"resume_id": resume["id"], "job_id": job["id"], "do_tailor": False},
    ).json()
    final = client.post(f"/tailoring/runs/{run['id']}/finalize", headers=AUTH).json()
    assert Path(final["final_tex_key"]).read_bytes() == original.encode()
