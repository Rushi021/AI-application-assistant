from __future__ import annotations

import os

os.environ["AGENT_WORKER_ENABLED"] = "false"
os.environ["SCRAPER_PROVIDER"] = "noop"
os.environ["LLM_DEFAULT"] = "fake:fake"
os.environ["LLM_JD_EXTRACTION"] = "fake:fake"
os.environ["LLM_RESUME_TAILOR"] = "fake:fake"
os.environ["LLM_BULLET_REWRITE"] = "fake:fake"
os.environ["LLM_COVER_LETTER"] = "fake:fake"
os.environ["EMBEDDING_DEFAULT"] = "fake:fake"
os.environ["EMBEDDING_REQUIREMENT"] = "fake:fake"
os.environ["EMBEDDING_RESUME"] = "fake:fake"

import pytest
from fastapi.testclient import TestClient

from app.modules.ai import providers as _providers  # noqa: F401


SAMPLE_TEX = r"""
\documentclass{article}
\begin{document}
\section{Experience}
\begin{itemize}
\item Built Python APIs for image classification on medical scans
\item Deployed Kubernetes services on AWS
\end{itemize}
\section{Skills}
Python, SQL, machine learning
\end{document}
"""

AUTH = {"Authorization": "Bearer test-token"}


def build_sample_docx() -> bytes:
    """A resume shaped like the real thing: caps headers, tab-aligned dates,
    bold lead-ins inside a line, an indented bullet list, and Education."""
    import io

    from docx import Document
    from docx.shared import Pt

    document = Document()

    name = document.add_paragraph()
    run = name.add_run("Rushi Patel")
    run.bold = True
    run.font.size = Pt(20)
    document.add_paragraph("rushi@example.com | 555-0100 | New York, NY")

    def header(text: str):
        paragraph = document.add_paragraph()
        header_run = paragraph.add_run(text)
        header_run.bold = True
        return paragraph

    def tabbed(left: str, right: str):
        paragraph = document.add_paragraph()
        left_run = paragraph.add_run(left)
        left_run.bold = True
        paragraph.add_run("\t")
        paragraph.add_run(right)
        return paragraph

    header("EDUCATION")
    tabbed("New York University", "May 2024")
    document.add_paragraph("M.S. Computer Science, GPA 3.8")

    header("EXPERIENCE")
    tabbed("Acme Corp - Data Scientist", "Jun 2024 - Present")
    for text in (
        "Built Python pipelines processing 2M records daily",
        "Trained image classification models on medical scans",
    ):
        bullet = document.add_paragraph(style="List Bullet")
        bullet.paragraph_format.left_indent = Pt(18)
        bullet.add_run(text)

    header("TECHNICAL SKILLS")
    for label, values in (("Languages: ", "Python, SQL, R"), ("Tools: ", "Docker, Kubernetes, AWS")):
        paragraph = document.add_paragraph()
        label_run = paragraph.add_run(label)
        label_run.bold = True
        paragraph.add_run(values)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def client(tmp_path, monkeypatch):
    db_path = tmp_path / "test.db"
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{db_path}")
    monkeypatch.setenv("LOCAL_STORAGE_PATH", str(tmp_path / "resumes"))
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path / "output"))
    monkeypatch.setenv("RESUME_LIBRARY_ROOT", str(tmp_path / "library"))
    # Never drive LibreOffice or Word from a test run.
    monkeypatch.setenv("DOCX_PDF_CONVERTER", "off")
    monkeypatch.setenv("LLM_DEFAULT", "fake:fake")
    monkeypatch.setenv("LLM_RESUME_TAILOR", "fake:fake")
    monkeypatch.setenv("EMBEDDING_DEFAULT", "fake:fake")

    from app.core.config import get_settings
    from app.modules.ai.gateway import reset_shared_gateway

    get_settings.cache_clear()
    reset_shared_gateway()

    from app.db.session import configure_engine
    from app.db.base import Base

    engine = configure_engine(f"sqlite:///{db_path}")
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)

    from app.main import create_app

    with TestClient(create_app()) as test_client:
        yield test_client
